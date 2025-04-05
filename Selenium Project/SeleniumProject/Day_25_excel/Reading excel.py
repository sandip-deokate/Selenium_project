#file --> workbook -->sheet --> Rows --> Cells
import openpyxl
import os
File_Path="C:\Arrow Incubator\BomShort.xlsx"
workbook = openpyxl.load_workbook(File_Path)
sheet= workbook["Best Buying Options"]

rows = sheet.max_row
cols = sheet.max_column

for r in range(1,rows+1):
    for c in range(1,cols+1):
        print(sheet.cell(r,c).value, end="    |")
    print("")




