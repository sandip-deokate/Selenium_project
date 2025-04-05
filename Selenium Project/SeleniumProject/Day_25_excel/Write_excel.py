import openpyxl

"""def ChromeClass():
    from selenium import webdriver
    from selenium.webdriver.chrome.service import Service
    from webdriver_manager.chrome import ChromeDriverManager
    from selenium.webdriver.chrome.options import Options
    ops= Options
    #ops.add_argument("start-maximized")

    driver = webdriver.Chrome(service=Service(ChromeDriverManager(driver_version="").install()),options=ops)
    return driver

driver = ChromeClass()"""

path = "C:\\Udemycourses\\GitandGitHub\\Selenium Project\\SeleniumProject\\Day_25_excel\\excel.xlsx"
workbook=openpyxl.load_workbook(path)
sheet=workbook.active

for i in range(1,6):
    for j in range(1,4):
        sheet.cell(i,j).value="welcome"

workbook.save(path)
for i in range(1,6):
    for j in range(1,4):
        print(sheet.cell(i,j).value,end=" ")
    print()