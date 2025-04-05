import openpyxl
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import PatternFill


def getRowCount(FilePath,sheetName):
     workbook = openpyxl.load_workbook(FilePath)
     sheet=workbook[sheetName]
     return(sheet.max_row)

def getColCount(FilePath,sheetName):
    workbook = openpyxl.load_workbook(FilePath)
    sheet = workbook[sheetName]
    return (sheet.max_column)

def readFile(FilePath,sheetName,rowNum,colNum):
    workbook = openpyxl.load_workbook(FilePath)
    sheet = workbook[sheetName]
    return sheet.cell(rowNum,colNum).value
def writeFile(FilePath,sheetName,rowNum,colNum,data):
    workbook = openpyxl.load_workbook(FilePath)
    sheet = workbook[sheetName]
    sheet.cell(rowNum,colNum).value= data
    workbook.save(FilePath)

def fillGreenColor(filePath,sheetName,rowNum,colNum):
    workbook =openpyxl.load_workbook(filePath)
    sheet=workbook[sheetName]
    greenFill= PatternFill(start_color='60b212',
                           end_color='60b212',
                           fill_type='solid')
    sheet.cell(rowNum,colNum).fill=greenFill
    workbook.save(filePath)

def fillRedColor(filePath,sheetName,rowNum,colNum):
    workbook = openpyxl.load_workbook(filePath)
    sheet = workbook[sheetName]
    redFill = PatternFill(start_color='ff0000',
                           end_color='ff0000',
                           fill_type='solid')
    sheet.cell(rowNum,colNum).fill=redFill
    workbook.save(filePath)











